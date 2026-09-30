"""图片描述并发化(滑动窗口)测试:md文档/DEV_concurrent_image_description.md §9。

fixture 由 python-docx 程序生成(8 张图,各自带唯一图注);模型一律假实现,
经模块级 describe_image patch 注入。并发行为用真实时钟与真实观测库验证。
"""

import io
import threading
import time
import tempfile
import unittest
from pathlib import Path

import _bootstrap  # noqa: F401

try:
    import numpy  # noqa: F401
    from PIL import Image
    import docx  # noqa: F401
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

if HAVE_DEPS:
    from deephoto import repo
    from deephoto.config import Settings
    from deephoto.db import connect, init_db
    from deephoto.indexing.service import IndexService
    from deephoto.parsing.formats import format_by_key
    from deephoto.pipeline import progress as pg
    from deephoto.pipeline.ingest import IngestService
    from deephoto.progress_store import ProgressStore
    from deephoto.security import LOCAL_CTX
    from deephoto.storage import ObjectStore

IMAGE_COUNT = 8


def _png(color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buf, format="PNG")
    return buf.getvalue()


def _docx_8_images() -> bytes:
    """8 张图,每张配唯一图注(图 1.i 独有说明i),供结果配对断言。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("正文" * 50)                    # 越过"正文几乎为空"回退阈值
    for i in range(1, IMAGE_COUNT + 1):
        p = doc.add_paragraph()
        p.add_run().add_picture(io.BytesIO(_png()))
        doc.add_paragraph(f"图 1.{i} 独有说明{i}", style="Caption")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _good(caption):
    return {"visible_summary": f"desc:{caption}", "visible_labels": [],
            "caption": caption or "", "context_summary": "",
            "uncertain_details": [], "diagnostic": {"result": "ok"}}


@unittest.skipUnless(HAVE_DEPS, "需要 numpy、pillow 与 python-docx")
class ConcurrentDescribeTest(unittest.TestCase):
    """并发度 4 的滑动窗口路径;观测库用真实时钟(耗时口径断言依赖真实时间)。"""

    CONCURRENCY = 4

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        root = Path(self.tmp.name)
        self.settings = Settings(
            moonshot_api_key=None, moonshot_base_url="", chat_model="k3", chat_temperature=1.0,
            embedding_base_url=None, embedding_api_key=None, embedding_model=None,
            data_dir=root, max_upload_mb=100, ingestion_version="v2", mineru_api_key=None,
            description_concurrency=self.CONCURRENCY)
        init_db(self.settings.db_path)
        self.conn = connect(self.settings.db_path)
        self.store = ObjectStore(root / "objects")
        self.prog = ProgressStore(root / "progress.db", instance_id="inst-conc")
        data = _docx_8_images()
        key, digest = self.store.put(data, "sources", format_by_key("docx").mime, ext="docx")
        self.doc_id = repo.insert_document(
            self.conn, tenant_id=LOCAL_CTX.tenant_id, owner_id=LOCAL_CTX.user_id,
            filename="八图.docx", source_object_key=key, sha256=digest, ingestion_version="v2",
            source_format="docx", locator_kind="section", parse_engine="local")
        self.prog.register_run(f"run_{self.doc_id}", self.doc_id, LOCAL_CTX.tenant_id)
        self.prog.claim_run(f"run_{self.doc_id}")

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def _ingest(self, describe, **overrides):
        """在调用线程内同步执行入库;describe 为假 describe_image(模块级 patch)。"""
        settings = self.settings
        if overrides:
            from dataclasses import replace
            settings = replace(self.settings, **overrides)
        service = IngestService(
            settings, self.store, parser=None, index_service=IndexService(),
            chat_model_factory=lambda: object(), progress_store=self.prog)
        import deephoto.pipeline.ingest as ingest_mod
        original = ingest_mod.describe_image
        ingest_mod.describe_image = describe
        try:
            service.ingest(self.doc_id)
        finally:
            ingest_mod.describe_image = original

    def _occs(self):
        return repo.occurrences_for_document(self.conn, self.doc_id)

    def _items(self):
        return self.prog.detail(LOCAL_CTX.tenant_id, self.doc_id)["items"]

    def _stage(self, name):
        stages = self.prog.detail(LOCAL_CTX.tenant_id, self.doc_id)["stages"]
        return {s["stage"]: s for s in stages}[name]

    # §9-1:并发确实发生(峰值 ≤ N 且 > 1),且窗口有界(§9-8 的计数面)
    def test_concurrency_peak_within_window(self):
        lock = threading.Lock()
        state = {"cur": 0, "peak": 0}

        def fake(model, data, mime, caption=None, **kw):
            with lock:
                state["cur"] += 1
                state["peak"] = max(state["peak"], state["cur"])
            time.sleep(0.05)
            with lock:
                state["cur"] -= 1
            return _good(caption)

        self._ingest(fake)
        self.assertEqual(repo.get_document(self.conn, self.doc_id)["status"], "ready")
        self.assertGreater(state["peak"], 1)                      # 确实并发了
        self.assertLessEqual(state["peak"], self.CONCURRENCY)     # 窗口有界

    # §9-2:并发度 1 走旧路径,不创建线程池
    def test_concurrency_1_never_builds_pool(self):
        import deephoto.pipeline.ingest as ingest_mod
        builds = []
        original = ingest_mod.ThreadPoolExecutor
        ingest_mod.ThreadPoolExecutor = lambda *a, **k: builds.append(1) or original(*a, **k)
        try:
            self._ingest(lambda model, data, mime, caption=None, **kw: _good(caption),
                         description_concurrency=1)
        finally:
            ingest_mod.ThreadPoolExecutor = original
        self.assertEqual(builds, [])
        self.assertEqual(len([o for o in self._occs() if o["description"]]), IMAGE_COUNT)

    # §9-3:结果落到正确的行,不串位
    def test_results_land_on_matching_occurrence(self):
        self._ingest(lambda model, data, mime, caption=None, **kw: _good(caption))
        occs = self._occs()
        self.assertEqual(len(occs), IMAGE_COUNT)
        for occ in occs:
            self.assertTrue(occ["caption"], occ)                       # 每张都有图注
            self.assertEqual(occ["description"], f"desc:{occ['caption']}")

    # §9-4:单图失败不影响其他;§9-5:格式失败单独口径
    def test_partial_failures_do_not_block_others(self):
        def fake(model, data, mime, caption=None, **kw):
            if "独有说明2" in (caption or ""):
                raise RuntimeError("boom-2")
            if "独有说明5" in (caption or ""):
                raise RuntimeError("boom-5")
            if "独有说明7" in (caption or ""):
                return {**_good(caption), "diagnostic": {"result": "parse_failed"}}
            return _good(caption)

        self._ingest(fake)
        doc = repo.get_document(self.conn, self.doc_id)
        self.assertEqual(doc["status"], "ready")                       # 整篇不失败
        stage = self._stage(pg.STAGE_DESCRIBING)
        self.assertEqual(stage["result"], pg.RESULT_PARTIAL)
        results = sorted(i["result"] for i in self._items() if i["kind"] == "image")
        self.assertEqual(results.count(pg.ITEM_OK), 5)
        self.assertEqual(results.count(pg.ITEM_ERROR), 2)
        self.assertEqual(results.count(pg.ITEM_PARSE_FAILED), 1)
        # 失败图降级为空描述,成功图描述完好
        for occ in self._occs():
            if "独有说明2" in occ["caption"] or "独有说明5" in occ["caption"]:
                self.assertEqual(occ["description"], "")
                self.assertTrue(occ["uncertain_details"])              # 失败原因留痕
            elif "独有说明7" not in occ["caption"]:
                self.assertTrue(occ["description"].startswith("desc:"))

    # §9-6:观测计数正确;§9-7:耗时口径(阶段耗时明显小于逐张之和,逐张不含排队)
    def test_counts_and_duration_semantics(self):
        def fake(model, data, mime, caption=None, **kw):
            time.sleep(0.05)
            return _good(caption)

        self._ingest(fake)
        stage = self._stage(pg.STAGE_DESCRIBING)
        counts = stage["counts"]
        self.assertEqual(counts["completed"], IMAGE_COUNT)
        self.assertEqual(counts["succeeded"], IMAGE_COUNT)
        items = [i for i in self._items() if i["kind"] == "image"]
        self.assertEqual(len(items), IMAGE_COUNT)
        durations = []
        for i in items:
            self.assertEqual(i["result"], pg.ITEM_OK)
            self.assertIsNotNone(i["duration_ms"])
            self.assertGreaterEqual(i["duration_ms"], 40)              # ≈ 单次调用时长
            durations.append(i["duration_ms"])
        # 并发 4、8 张:阶段耗时约 2 波,明显小于逐张耗时之和(各图耗时重叠,F12)
        self.assertLess(stage["duration_ms"], sum(durations) * 0.8)

    # §9-8:窗口有界(阻塞假模型下,已开始未结束的图数 ≤ N)
    def test_window_is_bounded_under_blocking_model(self):
        gate = threading.Event()
        entered = []
        lock = threading.Lock()

        def fake(model, data, mime, caption=None, **kw):
            with lock:
                entered.append(caption)
            if not gate.wait(10):
                raise RuntimeError("test gate timeout")
            return _good(caption)

        thread = threading.Thread(target=self._ingest, args=(fake,), daemon=True)
        thread.start()
        deadline = time.monotonic() + 10
        while len(entered) < self.CONCURRENCY and time.monotonic() < deadline:
            time.sleep(0.01)
        self.assertEqual(len(entered), self.CONCURRENCY)               # 只开了窗口大小
        running = self.prog._conn().execute(
            "SELECT COUNT(*) c FROM items WHERE result = 'running'").fetchone()["c"]
        self.assertEqual(running, self.CONCURRENCY)                    # 观测口径同样有界
        # §7-B:在途多于一张时列表行显示聚合文案。聚合标签在主线程 refresh 时写入,
        # 与假模型的 entered 计数存在微秒级先后,轮询等待其出现(闸门未开,不会翻回)
        want = f"{self.CONCURRENCY} 张并发处理中"
        label = None
        deadline = time.monotonic() + 10
        while label != want and time.monotonic() < deadline:
            s = self.prog.summaries(
                LOCAL_CTX.tenant_id, [{"id": self.doc_id, "status": "describing"}])[self.doc_id]
            label = (s["current_item"] or {}).get("label")
            time.sleep(0.01)
        self.assertEqual(label, want)
        gate.set()
        thread.join(30)
        self.assertFalse(thread.is_alive())
        self.assertEqual(len(entered), IMAGE_COUNT)
        self.assertEqual(repo.get_document(self.conn, self.doc_id)["status"], "ready")

    # §9-9:主线程中途异常 -> 池被关闭、未开始的任务被取消、文档按失败落库
    def test_main_thread_exception_cancels_pending_and_shuts_down(self):
        gate = threading.Event()
        entered = []
        lock = threading.Lock()

        def fake(model, data, mime, caption=None, **kw):
            with lock:
                entered.append(caption)
            if not gate.wait(10):
                raise RuntimeError("test gate timeout")
            return _good(caption)

        import deephoto.repo as repo_mod
        original = repo_mod.update_occurrence_description
        baseline = {t.name for t in threading.enumerate()}
        thread = threading.Thread(target=self._ingest, args=(fake,), daemon=True)
        thread.start()
        deadline = time.monotonic() + 10
        while len(entered) < self.CONCURRENCY and time.monotonic() < deadline:
            time.sleep(0.01)
        self.assertEqual(len(entered), self.CONCURRENCY)
        repo_mod.update_occurrence_description = lambda *a, **k: (_ for _ in ()).throw(
            RuntimeError("db boom"))
        try:
            gate.set()                       # 首批完成 -> 主线程 settle 时抛错
            thread.join(30)
        finally:
            repo_mod.update_occurrence_description = original
        self.assertFalse(thread.is_alive())
        self.assertEqual(repo.get_document(self.conn, self.doc_id)["status"], "failed")
        self.assertEqual(len(entered), self.CONCURRENCY)               # 未开始的不再提交
        deadline = time.monotonic() + 10                               # 池线程随即退出
        while time.monotonic() < deadline:
            pools = [t for t in threading.enumerate()
                     if t.name.startswith("ThreadPoolExecutor") and t.name not in baseline]
            if not pools:
                break
            time.sleep(0.05)
        self.assertEqual(pools, [])

    # §9-10:疑似限流识别与告警(不硬依赖 openai)
    def test_rate_limit_suspects_warn(self):
        class RateLimitError(Exception):
            pass

        def fake(model, data, mime, caption=None, **kw):
            if "独有说明1" in (caption or ""):
                exc = RuntimeError("too many requests")
                exc.status_code = 429
                raise exc
            if "独有说明4" in (caption or ""):
                raise RateLimitError("slow down")
            return _good(caption)

        self._ingest(fake)
        doc = repo.get_document(self.conn, self.doc_id)
        self.assertEqual(doc["status"], "ready")
        warnings = self.prog.detail(LOCAL_CTX.tenant_id, self.doc_id)["summary"]["warnings"]
        hits = [w for w in warnings if "DESCRIPTION_CONCURRENCY" in w]
        self.assertEqual(len(hits), 1, warnings)
        self.assertIn("2 张", hits[0])
        results = sorted(i["result"] for i in self._items() if i["kind"] == "image")
        self.assertEqual(results.count(pg.ITEM_ERROR), 2)
        self.assertEqual(results.count(pg.ITEM_OK), IMAGE_COUNT - 2)


if __name__ == "__main__":
    unittest.main()
