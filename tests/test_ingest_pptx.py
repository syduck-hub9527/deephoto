"""pptx 端到端入库:上传字段 -> 本地解析 -> 分块/图文关联 -> 索引 -> 检索(P2 §5 验收)。

不经网络、不经 MinerU;fixture 由 python-pptx 程序生成。
"""

import io
import tempfile
import unittest
from pathlib import Path

import _bootstrap  # noqa: F401

try:
    import numpy  # noqa: F401
    from PIL import Image
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

if HAVE_DEPS:
    from deephoto import repo
    from deephoto.agent.knowledge import KnowledgeService
    from deephoto.config import Settings
    from deephoto.db import connect, init_db
    from deephoto.indexing.service import IndexService
    from deephoto.pipeline.ingest import IngestService
    from deephoto.security import LOCAL_CTX
    from deephoto.storage import ObjectStore


def _png(color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buf, format="PNG")
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "存储器进展"
    s1.placeholders[1].text = "容量持续提升,如图 1.1 所示。"
    s1.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(3))
    cap = s1.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(0.5))
    cap.text_frame.text = "图 1.1 容量与尺寸"
    s1.notes_slide.notes_text_frame.text = "这一页重点讲容量趋势"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    gfx = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    gfx.table.rows[0].cells[0].text = "元素"
    gfx.table.rows[0].cells[1].text = "类型"
    gfx.table.rows[1].cells[0].text = "磷"
    gfx.table.rows[1].cells[1].text = "N 型"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@unittest.skipUnless(HAVE_DEPS, "需要 numpy、pillow 与 python-pptx")
class IngestPptxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        root = Path(self.tmp.name)
        self.settings = Settings(
            moonshot_api_key=None, moonshot_base_url="", chat_model="k3", chat_temperature=1.0,
            embedding_base_url=None, embedding_api_key=None, embedding_model=None,
            data_dir=root, max_upload_mb=100, ingestion_version="v2", mineru_api_key=None)
        init_db(self.settings.db_path)
        self.conn = connect(self.settings.db_path)
        self.store = ObjectStore(root / "objects")
        self.index = IndexService()
        self.doc_id = self._ingest(_pptx_bytes(), "讲义.pptx")
        self.knowledge = KnowledgeService(self.store, self.index)

    def _ingest(self, data: bytes, filename: str) -> str:
        from deephoto.parsing.formats import format_by_key
        key, digest = self.store.put(data, "sources", format_by_key("pptx").mime, ext="pptx")
        doc_id = repo.insert_document(
            self.conn, tenant_id=LOCAL_CTX.tenant_id, owner_id=LOCAL_CTX.user_id,
            filename=filename, source_object_key=key, sha256=digest, ingestion_version="v2",
            source_format="pptx", locator_kind="slide", parse_engine="local")
        IngestService(self.settings, self.store, parser=None,
                      index_service=self.index).ingest(doc_id)
        return doc_id

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def test_ready_with_slide_sections_and_notes(self):
        doc = repo.get_document(self.conn, self.doc_id)
        self.assertEqual(doc["status"], "ready")
        self.assertEqual(doc["page_count"], 2)                    # 一幻灯片一页
        chunks = repo.chunks_for_document(self.conn, self.doc_id)
        sections = {c["section"] for c in chunks}
        self.assertIn("幻灯片 1:存储器进展", sections)
        self.assertIn("幻灯片 2", sections)
        blob = "\n".join(c["text"] for c in chunks)
        self.assertIn("备注:这一页重点讲容量趋势", blob)           # 备注入块可检索
        self.assertIn("磷 | N 型", blob)                           # 表格文本可检索

    def test_figure_persisted_with_caption(self):
        occs = repo.occurrences_for_document(self.conn, self.doc_id)
        self.assertEqual(len(occs), 1)
        self.assertEqual(occs[0]["figure_number"], "1.1")
        self.assertEqual(occs[0]["page_number"], 1)
        asset = repo.get_asset(self.conn, occs[0]["image_asset_id"])
        self.assertEqual(asset["mime_type"], "image/png")
        self.assertTrue(self.store.get(asset["original_object_key"]).startswith(b"\x89PNG"))

    def test_search_gives_slide_locator(self):
        result = self.knowledge.search(self.conn, LOCAL_CTX, "容量 尺寸 存储器")
        self.assertTrue(result["chunks"])
        for c in result["chunks"]:
            self.assertTrue(c["locator"].startswith("幻灯片 "), c["locator"])
            self.assertEqual(c["document"], "讲义.pptx")
        images = result["images"]
        self.assertTrue(any(i["figure_number"] == "1.1" for i in images))
        for i in images:
            self.assertEqual(i["locator_label"], "幻灯片 1")

    def test_image_entry_has_no_page_preview(self):
        # 问答配图路径:非 PDF 文档 source_page_url 为 None(页预览仅 PDF,§3.9-4)
        occs = repo.occurrences_for_document(self.conn, self.doc_id)
        entries = self.knowledge.build_image_entries(
            self.conn, LOCAL_CTX, [o["id"] for o in occs])
        self.assertEqual(len(entries), 1)
        self.assertIsNone(entries[0]["source_page_url"])
        self.assertEqual(entries[0]["locator_label"], "幻灯片 1")

    def test_notes_slide_without_body_placeholder(self):
        # 回归:备注页缺正文占位符时 notes_text_frame 为 None,不能让整份 pptx 解析失败
        from pptx import Presentation
        from pptx.enum.shapes import PP_PLACEHOLDER
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "无备注占位符页"
        for ph in list(s.notes_slide.placeholders):
            if ph.placeholder_format.type == PP_PLACEHOLDER.BODY:
                ph._element.getparent().remove(ph._element)
        buf = io.BytesIO()
        prs.save(buf)
        doc_id = self._ingest(buf.getvalue(), "无占位符.pptx")
        self.assertEqual(repo.get_document(self.conn, doc_id)["status"], "ready")
        blob = "\n".join(c["text"] for c in repo.chunks_for_document(self.conn, doc_id))
        self.assertIn("无备注占位符页", blob)                      # 正文正常入库
        self.assertNotIn("备注:", blob)                            # 缺占位符不捏造备注

    def test_soft_break_does_not_leak_control_char(self):
        # 回归:强制换行(a:br)在 python-pptx 里是 \x0b,不能带进章节名与正文
        from pptx import Presentation
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text_frame.text = "主标题\x0b副标题"        # 写成真实 a:br
        s.placeholders[1].text_frame.text = "第一行\x0b第二行"
        buf = io.BytesIO()
        prs.save(buf)
        doc_id = self._ingest(buf.getvalue(), "换行标题.pptx")
        chunks = repo.chunks_for_document(self.conn, doc_id)
        self.assertIn("幻灯片 1:主标题 副标题", {c["section"] for c in chunks})
        blob = "\n".join(c["text"] for c in chunks)
        self.assertIn("第一行 第二行", blob)
        self.assertNotIn("\x0b", blob + "".join(c["section"] for c in chunks))


if __name__ == "__main__":
    unittest.main()
