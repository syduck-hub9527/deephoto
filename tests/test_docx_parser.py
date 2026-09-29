"""docx 本地解析器与 OOXML 安全闸门测试:fixture 全部由程序生成,不提交二进制样本(§8)。"""

import io
import unittest
import zipfile

import _bootstrap  # noqa: F401

try:
    import numpy  # noqa: F401
    from PIL import Image
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

from deephoto.config import Settings
from deephoto.parsing.formats import UnsupportedFormat, check_zip_safety, format_by_key
from deephoto.parsing.registry import EngineUnavailable, SourceFile, validate_parsed


def _settings(**overrides):
    kwargs = dict(
        moonshot_api_key=None, moonshot_base_url="", chat_model="k3", chat_temperature=1.0,
        embedding_base_url=None, embedding_api_key=None, embedding_model=None,
        data_dir="/tmp/x", max_upload_mb=100, ingestion_version="v2", mineru_api_key=None,
    )
    kwargs.update(overrides)
    return Settings(**kwargs)


class _Recorder:
    def __init__(self):
        self.warnings: list[str] = []

    def warn(self, message):
        self.warnings.append(message)


def _png(color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buf, format="PNG")
    return buf.getvalue()


def _docx(build) -> bytes:
    from docx import Document
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _parse(data: bytes, observer=None, **settings_over):
    from deephoto.parsing.docx_parser import DocxParser
    src = SourceFile(data=data, filename="a.docx", fmt=format_by_key("docx"))
    return DocxParser(_settings(**settings_over)).parse(src, observer or _Recorder())


@unittest.skipUnless(HAVE_DEPS, "需要 pillow 与 python-docx")
class DocxParserTest(unittest.TestCase):
    def test_headings_sections_and_paragraphs(self):
        data = _docx(lambda d: (
            d.add_heading("第1章 引论", level=1),
            d.add_paragraph("正文一段。"),
            d.add_heading("1.1 集成度", level=2),
            d.add_paragraph("如图 1.1 所示,容量持续增大。"),
        ))
        doc = _parse(data)
        texts = {p.text: p.section for p in doc.all_paragraphs()}
        self.assertEqual(texts["第1章 引论"], "第1章 引论")
        self.assertEqual(texts["如图 1.1 所示,容量持续增大。"], "第1章 引论 > 1.1 集成度")
        self.assertEqual(doc.locator_kind, "section")
        validate_parsed(doc)

    def test_wps_style_heading_fallback(self):
        # WPS 样式名可能是中文"标题 N"
        def build(d):
            d.styles.add_style("标题 2", 1)   # WD_STYLE_TYPE.PARAGRAPH
            d.add_paragraph("中文样式标题", style="标题 2")
            d.add_paragraph("正文。")
        doc = _parse(_docx(build))
        texts = {p.text: p.section for p in doc.all_paragraphs()}
        self.assertEqual(texts["正文。"], "中文样式标题")

    def test_style_definition_outline_level(self):
        # 回归:自定义样式的大纲级别写在样式定义里(w:style/w:pPr/w:outlineLvl),
        # 段落自身 pPr 没有,修前不当标题、section 为空
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        def _set_outline(style, val: str):
            ppr = style.element.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                style.element.append(ppr)
            ol = OxmlElement("w:outlineLvl")
            ol.set(qn("w:val"), val)
            ppr.append(ol)

        def build(d):
            parent = d.styles.add_style("我的标题", 1)   # WD_STYLE_TYPE.PARAGRAPH
            _set_outline(parent, "0")
            child = d.styles.add_style("我的子标题", 1)
            bo = OxmlElement("w:basedOn")
            bo.set(qn("w:val"), parent.style_id)
            child.element.append(bo)                   # 子样式不写 outlineLvl,沿 basedOn 继承
            d.add_paragraph("自定义样式标题", style="我的标题")
            d.add_paragraph("正文一。")
            d.add_paragraph("继承的子标题", style="我的子标题")
            d.add_paragraph("正文二。")

        doc = _parse(_docx(build))
        texts = {p.text: p.section for p in doc.all_paragraphs()}
        self.assertEqual(texts["正文一。"], "自定义样式标题")
        # 子样式不写 outlineLvl,沿 basedOn 继承为 1 级;同级标题替换而非嵌套
        self.assertEqual(texts["正文二。"], "继承的子标题")

    def test_image_with_caption_after(self):
        def build(d):
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不会再回退到云端解析,这一段是填充。")
            d.add_paragraph("示意图如下:")
            p = d.add_paragraph()
            p.add_run().add_picture(io.BytesIO(_png()))
            d.add_paragraph("图 2.1 容量与尺寸", style="Caption")
            d.add_paragraph("后续正文。")
        doc = _parse(_docx(build))
        figs = doc.all_figures()
        self.assertEqual(len(figs), 1)
        caps = doc.all_captions()
        self.assertEqual(len(caps), 1)
        self.assertEqual(caps[0].figure_number, "2.1")
        self.assertEqual(figs[0].caption_id, caps[0].id)
        # 图注文本落在段落里(caption_of 关联的前提)
        self.assertIn("图 2.1 容量与尺寸", [p.text for p in doc.all_paragraphs()])
        validate_parsed(doc)

    def test_image_with_caption_before(self):
        def build(d):
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不会再回退到云端解析,这一段是填充。")
            d.add_paragraph("图 1.1 前置图注", style="Caption")
            p = d.add_paragraph()
            p.add_run().add_picture(io.BytesIO(_png("blue")))
        doc = _parse(_docx(build))
        figs = doc.all_figures()
        self.assertEqual(len(figs), 1)
        self.assertEqual(doc.all_captions()[0].figure_number, "1.1")
        # 配对成功的前置图注只出现一次(图片的图注段落),不再单独出文本块
        self.assertEqual([p.text for p in doc.all_paragraphs()].count("图 1.1 前置图注"), 1)

    def test_caption_after_image_not_stolen_by_next_figure(self):
        # 回归:图1→"图1.1 第一张图"→图2→"图1.2 第二张图",修前两张图都配到"图1.1",
        # 图注文本重复出现在 3 个段落里
        filler = "这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不会再回退到云端解析,这一段是填充。"

        def build(d):
            d.add_paragraph(filler)
            d.add_paragraph().add_run().add_picture(io.BytesIO(_png("red")))
            d.add_paragraph("图 1.1 第一张图", style="Caption")
            d.add_paragraph().add_run().add_picture(io.BytesIO(_png("blue")))
            d.add_paragraph("图 1.2 第二张图", style="Caption")
        doc = _parse(_docx(build))
        figs = doc.all_figures()
        self.assertEqual(len(figs), 2)
        caps = {c.id: c for c in doc.all_captions()}
        self.assertEqual(caps[figs[0].caption_id].figure_number, "1.1")   # 各配各的
        self.assertEqual(caps[figs[1].caption_id].figure_number, "1.2")
        texts = [p.text for p in doc.all_paragraphs()]
        self.assertEqual(texts.count("图 1.1 第一张图"), 1)   # 图注文本恰好一份(装配层图注段落)
        self.assertEqual(texts.count("图 1.2 第二张图"), 1)

    def test_unpaired_caption_falls_back_to_plain_text(self):
        # 附近没有图的图注样式段(如表格标题):不丢,落为普通文本恰好一份
        def build(d):
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不回退。")
            d.add_paragraph("表 1.1 参数对照", style="Caption")
            d.add_paragraph("表格见上。")
        doc = _parse(_docx(build))
        texts = [p.text for p in doc.all_paragraphs()]
        self.assertEqual(texts.count("表 1.1 参数对照"), 1)
        self.assertEqual(doc.all_figures(), [])

    def test_content_controls_unwrapped(self):
        # 回归:块级 w:sdt / w:customXml 包住的段落与表格修前整段丢失(目录/封面/模板常见)
        from docx.oxml import OxmlElement
        filler = "这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不回退。"

        def build(d):
            d.add_paragraph(filler)
            inner = d.add_paragraph("控件里的正文")
            tbl_para = d.add_paragraph("控件表前的说明")
            t = d.add_table(rows=2, cols=1)
            t.rows[0].cells[0].text = "控件里的表头"
            t.rows[1].cells[0].text = "控件里的单元格"
            custom = d.add_paragraph("customXml 里的正文")
            # XML 手术:把控件段落/表格包进 w:sdt > w:sdtContent;另一段包进 w:customXml
            body = d.element.body
            for el in (inner._p, tbl_para._p, t._tbl):
                body.remove(el)
            sdt = OxmlElement("w:sdt")
            content = OxmlElement("w:sdtContent")
            for el in (inner._p, tbl_para._p, t._tbl):
                content.append(el)
            sdt.append(content)
            body.append(sdt)
            body.remove(custom._p)
            cx = OxmlElement("w:customXml")
            cx.append(custom._p)
            body.append(cx)

        doc = _parse(_docx(build))
        texts = [p.text for p in doc.all_paragraphs()]
        self.assertIn("控件里的正文", texts)
        self.assertIn("customXml 里的正文", texts)
        blob = "\n".join(texts)
        self.assertIn("控件里的单元格", blob)      # 容器里的表格也在
        validate_parsed(doc)

    def test_tracked_changes_accepted_view(self):
        # 回归:开着修订的文档,w:ins 新增文字修前整段丢失;w:del 删除内容不应出现
        from docx.oxml import OxmlElement

        def build(d):
            p = d.add_paragraph()
            p.add_run("原文一二三")
            ins = OxmlElement("w:ins")
            r = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = "修订新增内容XYZ"
            r.append(t)
            ins.append(r)
            p._p.append(ins)
            # w:del:删除内容(接受后视角)不得出现
            p2 = d.add_paragraph()
            p2.add_run("保留的文字")
            dele = OxmlElement("w:del")
            r2 = OxmlElement("w:r")
            dt = OxmlElement("w:delText")
            dt.text = "被删除的内容"
            r2.append(dt)
            dele.append(r2)
            p2._p.append(dele)
            # 表格单元格里的修订同样生效
            tb = d.add_table(rows=1, cols=1)
            cell_p = tb.rows[0].cells[0].paragraphs[0]
            cell_p.add_run("单元格原文")
            ins2 = OxmlElement("w:ins")
            r3 = OxmlElement("w:r")
            t3 = OxmlElement("w:t")
            t3.text = "单元格新增"
            r3.append(t3)
            ins2.append(r3)
            cell_p._p.append(ins2)

        doc = _parse(_docx(build))
        blob = "\n".join(p.text for p in doc.all_paragraphs())
        self.assertIn("原文一二三", blob)
        self.assertIn("修订新增内容XYZ", blob)      # w:ins 收入
        self.assertIn("保留的文字", blob)
        self.assertNotIn("被删除的内容", blob)      # w:del 跳过
        self.assertIn("单元格原文", blob)
        self.assertIn("单元格新增", blob)           # 表格单元格同规则

    def test_nested_table_and_cell_image(self):
        # 回归:嵌套表格的内层文字、单元格里的图片,修前都丢
        def build(d):
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不回退。")
            t = d.add_table(rows=1, cols=1)
            cell = t.rows[0].cells[0]
            cell.paragraphs[0].text = "外层单元格"
            nested = cell.add_table(rows=1, cols=1)
            nested.rows[0].cells[0].text = "内层嵌套文字"
            cell2 = cell.add_paragraph("说明 ")
            cell2.add_run().add_picture(io.BytesIO(_png("green")))

        doc = _parse(_docx(build))
        blob = "\n".join(p.text for p in doc.all_paragraphs())
        self.assertIn("外层单元格", blob)
        self.assertIn("内层嵌套文字", blob)        # 嵌套表格展平
        self.assertIn("说明", blob)
        self.assertEqual(len(doc.all_figures()), 1)   # 单元格图片成为图块
        validate_parsed(doc)

    def test_merged_cells_deduped(self):
        # 回归:横向合并单元格,row.cells 重复返回同一单元格,文字重复输出三次
        def build(d):
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不回退。")
            t = d.add_table(rows=2, cols=3)
            merged = t.rows[0].cells[0].merge(t.rows[0].cells[1]).merge(t.rows[0].cells[2])
            merged.text = "总标题合并"
            t.rows[1].cells[0].text = "甲"
            t.rows[1].cells[1].text = "乙"
            t.rows[1].cells[2].text = "丙"
        doc = _parse(_docx(build))
        blob = "\n".join(p.text for p in doc.all_paragraphs())
        self.assertEqual(blob.count("总标题合并"), 1)          # 合并文字只出现一次
        self.assertIn("甲 | 乙 | 丙", blob)                    # 普通行不受影响

    def test_table_rows_with_header_repeat_on_continuation(self):
        def build(d):
            d.add_paragraph("参数表:")     # 充字数,避免触发回退
            d.add_paragraph("正文填充,确保总字数超过五十个字的阈值,不回退。")
            t = d.add_table(rows=61, cols=2)
            t.rows[0].cells[0].text = "元素"
            t.rows[0].cells[1].text = "类型"
            for i in range(1, 61):
                t.rows[i].cells[0].text = f"元素{i}"
                t.rows[i].cells[1].text = "类型类型类型填充填充填充填充填充填充"
        doc = _parse(_docx(build))
        texts = [p.text for p in doc.all_paragraphs()]
        table_pieces = [t for t in texts if t.startswith("元素 | 类型")]
        self.assertGreater(len(table_pieces), 1)   # 续段重复表头(_table_paragraphs 约定)
        self.assertTrue(any("元素60" in t for t in texts))

    def test_emf_image_skipped_with_warning(self):
        # zip 级手术:把内嵌 PNG 改名为 .emf 并改 ContentType,模拟 EMF 图
        data = _docx(lambda d: (
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不会再回退到云端解析,这一段是填充。"),
            d.add_paragraph().add_run().add_picture(io.BytesIO(_png())),
        ))
        buf = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(buf, "w") as zout:
            for info in zin.infolist():
                payload = zin.read(info.filename)
                name = info.filename
                if name.startswith("word/media/"):
                    name = name.rsplit(".", 1)[0] + ".emf"
                if name == "[Content_Types].xml":
                    payload = payload.replace(b'Extension="png" ContentType="image/png"',
                                              b'Extension="emf" ContentType="image/x-emf"')
                if name == "word/_rels/document.xml.rels":
                    import re as _re
                    payload = _re.sub(rb'media/([\w.-]+)\.png', rb'media/\1.emf', payload)
                zout.writestr(name, payload)
        observer = _Recorder()
        doc = _parse(buf.getvalue(), observer)
        self.assertEqual(doc.all_figures(), [])
        self.assertTrue(any("跳过暂不支持的图片格式" in w for w in observer.warnings))

    def test_entity_probe_not_expanded(self):
        # 外部实体探针:python-docx 1.2.0 resolve_entities=False,实体不展开
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("[Content_Types].xml",
                       '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                       '<Default Extension="xml" ContentType="application/xml"/>'
                       '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            z.writestr("_rels/.rels",
                       '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            z.writestr("word/document.xml",
                       '<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY xxe "EXPANDED-SECRET">]>'
                       '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                       "<w:body><w:p><w:r><w:t>正文 &xxe; 结束,这段正文写得足够长,用来越过五十个字的回退阈值线。</w:t></w:r></w:p></w:body></w:document>")
        doc = _parse(buf.getvalue())
        blob = "\n".join(p.text for p in doc.all_paragraphs())
        self.assertNotIn("EXPANDED-SECRET", blob)
        self.assertIn("正文", blob)

    def test_fallback_to_mineru_needs_token_readable_error(self):
        # 只有图没有字:内容可能在文本框/形状里;无 Token 时报可读错误(不静默)
        data = _docx(lambda d: d.add_paragraph().add_run().add_picture(io.BytesIO(_png())))
        with self.assertRaisesRegex(EngineUnavailable, "DEEPHOTO_MINERU_API_KEY"):
            _parse(data)

    def test_image_rich_doc_with_text_does_not_fallback(self):
        def build(d):
            d.add_paragraph("这份文档的正文内容足够长,超过了五十个字的回退阈值,本地解析可以正常进行,不会再回退到云端解析,这一段是填充。")
            d.add_paragraph().add_run().add_picture(io.BytesIO(_png()))
        doc = _parse(_docx(build))   # 无 Token 也不回退
        self.assertEqual(len(doc.all_figures()), 1)


class ZipSafetyTest(unittest.TestCase):
    def _zip(self, entries: dict) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, payload in entries.items():
                z.writestr(name, payload)
        return buf.getvalue()

    def test_normal_zip_ok(self):
        check_zip_safety(self._zip({"word/document.xml": b"x" * 100}), 500)

    def test_entry_count_limit(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for i in range(10001):
                z.writestr(f"f{i}.txt", "x")
        with self.assertRaisesRegex(UnsupportedFormat, "条目数"):
            check_zip_safety(buf.getvalue(), 500)

    def test_path_traversal_rejected(self):
        for name in ("../evil.xml", "/abs/path.xml", "C:/abs.xml"):
            with self.assertRaises(UnsupportedFormat, msg=name):
                check_zip_safety(self._zip({name: b"x"}), 500)

    def test_compression_ratio_rejected(self):
        with self.assertRaisesRegex(UnsupportedFormat, "压缩比"):
            check_zip_safety(self._zip({"big.xml": b"0" * 5_000_000}), 500)

    def test_total_size_limit(self):
        import os
        with self.assertRaisesRegex(UnsupportedFormat, "解压总量"):
            check_zip_safety(self._zip({"a.xml": os.urandom(700_000), "b.xml": os.urandom(700_000)}), 1)


if __name__ == "__main__":
    unittest.main()
