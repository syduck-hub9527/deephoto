"""图片描述 parse_failed 集中重试测试:md文档/DEV_describe_parse_retry.md §8。

fixture 与 test_describe_concurrency.py 同风格(程序生成 docx、真实观测库);
模型一律假实现,经模块级 describe_image patch 注入;轮间退避 sleep 注入空函数。
"""

import io
import threading
import time
import tempfile
import unittest
from pathlib import Path

import _bootstrap  # noqa: F401

try:
    import numpy  # noqa: F401
    from PIL import Image
    import docx  # noqa: F401
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

if HAVE_DEPS:
    from deephoto import repo
    from deephoto.config import Settings
    from deephoto.db import connect, init_db
    from deephoto.indexing.service import IndexService
    from deephoto.parsing.formats import format_by_key
    from deephoto.pipeline import progress as pg
    from deephoto.pipeline.describe import DESC_PROMPT_VERSION, parse_description_json
    from deephoto.pipeline.ingest import IngestService
    from deephoto.progress_store import ProgressStore
    from deephoto.security import LOCAL_CTX
    from deephoto.storage import ObjectStore

IMAGE_COUNT = 8


def _png(color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buf, format="PNG")
    return buf.getvalue()


def _docx_8_images() -> bytes:
    """8 张图,每张配唯一图注(图 1.i 独有说明i),供结果配对断言。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("正文" * 50)                    # 越过"正文几乎为空"回退阈值
    for i in range(1, IMAGE_COUNT + 1):
        p = doc.add_paragraph()
        p.add_run().add_picture(io.BytesIO(_png()))
        doc.add_paragraph(f"图 1.{i} 独有说明{i}", style="Caption")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _good(caption):
    return {"visible_summary": f"desc:{caption}", "visible_labels": [],
            "caption": caption or "", "context_summary": "",
            "uncertain_details": [], "diagnostic": {"result": "ok"}}


def _parse_failed(reason="模型未返回 JSON 对象", code="no_json", truncated=False):
    """格式失败的假返回:diagnostic 带机读字段(与 describe.py §4.5 结构一致)。"""
    return {"visible_summary": "", "visible_labels": [], "caption": "",
            "context_summary": "", "uncertain_details": [reason],
            "diagnostic": {"result": "parse_failed", "reason": reason,
                           "code": code, "truncated": truncated}}


@unittest.skipUnless(HAVE_DEPS, "需要 numpy、pillow 与 python-docx")
class DescribeRetryTest(unittest.TestCase):
    """默认:重试 2 轮、并发 1(顺序路径)、注入空 sleep 与可记录的重试模型工厂。"""

    CONCURRENCY = 1

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        root = Path(self.tmp.name)
        self.settings = Settings(
            moonshot_api_key=None, moonshot_base_url="", chat_model="k3", chat_temperature=1.0,
            embedding_base_url=None, embedding_api_key=None, embedding_model=None,
            data_dir=root, max_upload_mb=100, ingestion_version="v2", mineru_api_key=None,
            description_concurrency=self.CONCURRENCY,
            description_parse_retry_rounds=2)
        init_db(self.settings.db_path)
        self.conn = connect(self.settings.db_path)
        self.store = ObjectStore(root / "objects")
        self.prog = ProgressStore(root / "progress.db", instance_id="inst-retry")
        data = _docx_8_images()
        key, digest = self.store.put(data, "sources", format_by_key("docx").mime, ext="docx")
        self.doc_id = repo.insert_document(
            self.conn, tenant_id=LOCAL_CTX.tenant_id, owner_id=LOCAL_CTX.user_id,
            filename="八图.docx", source_object_key=key, sha256=digest, ingestion_version="v2",
            source_format="docx", locator_kind="section", parse_engine="local")
        self.prog.register_run(f"run_{self.doc_id}", self.doc_id, LOCAL_CTX.tenant_id)
        self.prog.claim_run(f"run_{self.doc_id}")
        self.base_model = object()                     # 首轮模型(身份断言用)
        self.factory_calls = []                        # retry_model_factory 收到的 max_tokens

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def _factory(self, mt):
        self.factory_calls.append(mt)
        return ("retry-model", mt)                     # 每个目标上限一个可区分实例

    def _ingest(self, describe, with_factory=True, index_service=None, **overrides):
        """在调用线程内同步执行入库;describe 为假 describe_image(模块级 patch)。"""
        from dataclasses import replace
        settings = replace(self.settings, **overrides) if overrides else self.settings
        service = IngestService(
            settings, self.store, parser=None,
            index_service=index_service if index_service is not None else IndexService(),
            chat_model_factory=lambda: self.base_model,
            progress_store=self.prog,
            retry_model_factory=self._factory if with_factory else None,
            sleep=lambda _seconds: None)                # 轮间退避不等真实时钟
        import deephoto.pipeline.ingest as ingest_mod
        original = ingest_mod.describe_image
        ingest_mod.describe_image = describe
        try:
            service.ingest(self.doc_id)
        finally:
            ingest_mod.describe_image = original

    def _occs(self):
        return repo.occurrences_for_document(self.conn, self.doc_id)

    def _items(self, kind=None):
        items = self.prog.detail(LOCAL_CTX.tenant_id, self.doc_id)["items"]
        return [i for i in items if kind is None or i["kind"] == kind]

    def _stage(self, name):
        stages = self.prog.detail(LOCAL_CTX.tenant_id, self.doc_id)["stages"]
        return {s["stage"]: s for s in stages}.get(name)

    def _warnings(self):
        return self.prog.detail(LOCAL_CTX.tenant_id, self.doc_id)["summary"]["warnings"]

    def _occ_by_caption(self, needle):
        return next(o for o in self._occs() if needle in (o["caption"] or ""))

    # §8-1:全部 ok -> 调用次数 = 图数;无 describing_retry 子阶段;无重试 warn
    def test_all_ok_never_retries(self):
        calls = []

        def fake(model, data, mime, caption=None, **kw):
            calls.append(caption)
            return _good(caption)

        self._ingest(fake)
        self.assertEqual(len(calls), IMAGE_COUNT)
        self.assertIsNone(self._stage(pg.STAGE_DESCRIBE_RETRY))
        self.assertEqual(self._items("image_retry"), [])
        self.assertEqual(self._warnings(), [])
        self.assertEqual(self.factory_calls, [])

    # §8-2:1 张首轮 parse_failed、第 1 轮重试 ok -> 覆盖描述;无 warn;首轮与重试记录并存
    def test_recover_on_first_retry_round(self):
        calls = []

        def fake(model, data, mime, caption=None, retry_hint=None, **kw):
            calls.append((caption, retry_hint))
            if "独有说明3" in (caption or "") and retry_hint is None:
                return _parse_failed()
            return _good(caption)

        self._ingest(fake)
        doc = repo.get_document(self.conn, self.doc_id)
        self.assertEqual(doc["status"], "ready")
        # 描述被覆盖为重试结果,模型标记带轮次(§4.6)
        occ = self._occ_by_caption("独有说明3")
        self.assertEqual(occ["description"], "desc:图 1.3 独有说明3")
        self.assertTrue(occ["description_model"].endswith("+r1"), occ["description_model"])
        self.assertIn(DESC_PROMPT_VERSION, occ["description_model"])
        # 重试调用带了含失败原因的提示(§4.4)
        retry_hints = [hint for _caption, hint in calls if hint is not None]
        self.assertEqual(len(retry_hints), 1)
        self.assertIn("上一次输出无法解析为 JSON", retry_hints[0])
        self.assertIn("模型未返回 JSON 对象", retry_hints[0])
        # 统计修正后 describing 阶段 succeeded(§4.7.4)
        stage = self._stage(pg.STAGE_DESCRIBING)
        self.assertEqual(stage["result"], pg.RESULT_SUCCEEDED)
        # 首轮 parse_failed 与重试 ok 两条记录并存(§4.7.2);重试单项带轮次与 max_tokens
        first = self._items("image")
        self.assertEqual(sorted(i["result"] for i in first).count(pg.ITEM_PARSE_FAILED), 1)
        retries = self._items("image_retry")
        self.assertEqual(len(retries), 1)
        self.assertEqual(retries[0]["result"], pg.ITEM_OK)
        self.assertEqual(retries[0]["detail"]["round"], 1)
        self.assertEqual(retries[0]["detail"]["max_tokens"], 1024)   # no_json 不升级(§4.4)
        # 子阶段计数(§4.7.1);恢复后无任何 warn
        counts = self._stage(pg.STAGE_DESCRIBE_RETRY)["counts"]
        self.assertEqual(counts, {"total": 1, "recovered": 1, "still_failed": 0, "rounds": 1})
        self.assertEqual(self._warnings(), [])

    # §8-3:持续 parse_failed -> 调用 1+轮数 次;保持降级;仅 1 条最终 warn;阶段 partial
    def test_persistent_parse_failed_stays_degraded(self):
        calls = []

        def fake(model, data, mime, caption=None, retry_hint=None, **kw):
            calls.append(caption)
            if "独有说明6" in (caption or ""):
                return _parse_failed()
            return _good(caption)

        self._ingest(fake)
        self.assertEqual(calls.count("图 1.6 独有说明6"), 1 + 2)       # 首轮 + 2 轮重试
        doc = repo.get_document(self.conn, self.doc_id)
        self.assertEqual(doc["status"], "ready")                       # 整篇不失败
        occ = self._occ_by_caption("独有说明6")
        self.assertEqual(occ["description"], "")                       # 保持降级
        self.assertTrue(any("已重试 2 次仍失败" in d for d in occ["uncertain_details"]),
                        occ["uncertain_details"])
        self.assertEqual(self._stage(pg.STAGE_DESCRIBING)["result"], pg.RESULT_PARTIAL)
        retry_stage = self._stage(pg.STAGE_DESCRIBE_RETRY)
        self.assertEqual(retry_stage["result"], pg.RESULT_PARTIAL)
        self.assertEqual(retry_stage["counts"],
                         {"total": 1, "recovered": 0, "still_failed": 1, "rounds": 2})
        warnings = self._warnings()
        self.assertEqual(len(warnings), 1, warnings)                   # 首轮不 warn,仅最终 1 条
        self.assertIn("第 6 张", warnings[0])
        self.assertIn("已重试 2 次", warnings[0])
        self.assertIn("已降级", warnings[0])
        # image_retry 单项最终为 parse_failed(seq 沿用原图序号,§4.7.2)
        retries = self._items("image_retry")
        self.assertEqual([(i["seq"], i["result"]) for i in retries],
                         [(6, pg.ITEM_PARSE_FAILED)])

    # §8-4:轮数 = 0 -> 与旧行为一致:首轮立即 warn,无重试调用,无子阶段
    def test_rounds_zero_keeps_legacy_behavior(self):
        calls = []

        def fake(model, data, mime, caption=None, **kw):
            calls.append(caption)
            if "独有说明2" in (caption or ""):
                return _parse_failed()
            return _good(caption)

        self._ingest(fake, description_parse_retry_rounds=0)
        self.assertEqual(len(calls), IMAGE_COUNT)                      # 无重试调用
        self.assertIsNone(self._stage(pg.STAGE_DESCRIBE_RETRY))
        self.assertEqual(self._items("image_retry"), [])
        self.assertEqual(self.factory_calls, [])
        warnings = self._warnings()
        self.assertEqual(len(warnings), 1, warnings)
        self.assertIn("第 2 张", warnings[0])
        self.assertIn("已降级", warnings[0])
        self.assertNotIn("已重试", warnings[0])                        # 旧文案无重试字样

    # §8-5:error 项不进入重试清单;parse_failed 项正常重试
    def test_error_items_are_not_retried(self):
        calls = []

        def fake(model, data, mime, caption=None, **kw):
            calls.append(caption)
            if "独有说明2" in (caption or ""):
                raise RuntimeError("boom-2")
            if "独有说明5" in (caption or ""):
                return _parse_failed()
            return _good(caption)

        self._ingest(fake, description_parse_retry_rounds=1)
        self.assertEqual(calls.count("图 1.2 独有说明2"), 1)           # 异常项不重试
        self.assertEqual(calls.count("图 1.5 独有说明5"), 2)           # 格式失败项重试 1 轮
        self.assertEqual(repo.get_document(self.conn, self.doc_id)["status"], "ready")
        retry_seqs = sorted(i["seq"] for i in self._items("image_retry"))
        self.assertEqual(retry_seqs, [5])                              # 只有第 5 张进过重试
        counts = self._stage(pg.STAGE_DESCRIBE_RETRY)["counts"]
        self.assertEqual(counts["total"], 1)

    # §8-6:截断 -> max_tokens 逐轮翻倍并封顶;同一上限只建一次模型(§4.4)
    def test_truncation_escalates_max_tokens(self):
        def fake(model, data, mime, caption=None, **kw):
            if "独有说明4" in (caption or ""):
                return _parse_failed(reason="JSON 解析失败:JSONDecodeError",
                                     code="invalid_json", truncated=True)
            return _good(caption)

        self._ingest(fake)
        self.assertEqual(self.factory_calls, [2048, 4096])             # 1024 起翻倍,封顶 4096
        retries = self._items("image_retry")
        self.assertEqual(retries[0]["detail"]["max_tokens"], 4096)     # 末轮记录最终上限

    # §8-7:未装配 retry_model_factory -> 复用首轮模型,仍带重试提示,不报错
    def test_without_factory_reuses_first_pass_model(self):
        seen = []

        def fake(model, data, mime, caption=None, retry_hint=None, **kw):
            seen.append((model, retry_hint))
            if "独有说明7" in (caption or "") and retry_hint is None:
                return _parse_failed()
            return _good(caption)

        self._ingest(fake, with_factory=False, description_parse_retry_rounds=1)
        retry_seen = [(m, h) for m, h in seen if h is not None]
        self.assertEqual(len(retry_seen), 1)
        self.assertIs(retry_seen[0][0], self.base_model)               # 复用首轮模型
        self.assertIn("上一次输出无法解析为 JSON", retry_seen[0][1])   # 提示仍生效
        self.assertEqual(self._occ_by_caption("独有说明7")["description"],
                         "desc:图 1.7 独有说明7")

    # §8-8:并发路径同样成立;重试轮同一时刻在途数不超过并发度
    def test_concurrent_path_retries_within_window(self):
        lock = threading.Lock()
        state = {"cur": 0, "peak": 0}
        failing = {"图 1.2 独有说明2", "图 1.4 独有说明4", "图 1.6 独有说明6"}

        def fake(model, data, mime, caption=None, retry_hint=None, **kw):
            if caption in failing and retry_hint is None:
                return _parse_failed()
            if retry_hint is not None:                                 # 只统计重试轮在途
                with lock:
                    state["cur"] += 1
                    state["peak"] = max(state["peak"], state["cur"])
                time.sleep(0.05)
                with lock:
                    state["cur"] -= 1
            return _good(caption)

        self._ingest(fake, description_concurrency=4)
        for needle in failing:
            self.assertEqual(self._occ_by_caption(needle)["description"],
                             f"desc:{needle}")
        self.assertGreater(state["peak"], 1)                           # 重试轮确实并发
        self.assertLessEqual(state["peak"], 4)                         # 窗口有界
        counts = self._stage(pg.STAGE_DESCRIBE_RETRY)["counts"]
        self.assertEqual(counts, {"total": 3, "recovered": 3, "still_failed": 0, "rounds": 1})
        self.assertEqual(self._stage(pg.STAGE_DESCRIBING)["result"], pg.RESULT_SUCCEEDED)

    # §8-9:重试中抛异常 -> 文档仍 ready;该项记为仍失败;子阶段不向上抛
    def test_retry_exception_never_fails_document(self):
        def fake(model, data, mime, caption=None, retry_hint=None, **kw):
            if "独有说明8" in (caption or ""):
                if retry_hint is None:
                    return _parse_failed()
                raise RuntimeError("retry boom")                       # 重试轮异常
            return _good(caption)

        self._ingest(fake, description_parse_retry_rounds=1)
        self.assertEqual(repo.get_document(self.conn, self.doc_id)["status"], "ready")
        occ = self._occ_by_caption("独有说明8")
        self.assertEqual(occ["description"], "")
        self.assertTrue(any("已重试 1 次仍失败" in d for d in occ["uncertain_details"]),
                        occ["uncertain_details"])
        retries = self._items("image_retry")
        self.assertEqual(retries[0]["result"], pg.ITEM_ERROR)
        self.assertEqual(self._stage(pg.STAGE_DESCRIBING)["result"], pg.RESULT_PARTIAL)

    # §8-10:顺序保证——重试全部结束后才建索引,索引看到的是重试后的描述
    def test_index_built_after_retry_with_retried_description(self):
        seen_at_upsert = {}

        class SpyIndex(IndexService):
            def upsert_document(self2, conn, document_id, observer=None):
                row = conn.execute(
                    "SELECT description FROM image_occurrences WHERE document_id = ?",
                    (document_id,)).fetchall()
                seen_at_upsert["descriptions"] = sorted(r["description"] for r in row)
                return super().upsert_document(conn, document_id, observer=observer)

        def fake(model, data, mime, caption=None, retry_hint=None, **kw):
            if "独有说明1" in (caption or "") and retry_hint is None:
                return _parse_failed()
            return _good(caption)

        self._ingest(fake, index_service=SpyIndex(), description_parse_retry_rounds=1)
        descriptions = seen_at_upsert["descriptions"]
        self.assertEqual(len([d for d in descriptions if d]), IMAGE_COUNT)
        self.assertIn("desc:图 1.1 独有说明1", descriptions)           # 索引用重试后描述

    # §8-3 补充:轮间退避按注入的 sleep 执行(轮数 - 1 次)
    def test_backoff_between_rounds_uses_injected_sleep(self):
        sleeps = []

        def fake(model, data, mime, caption=None, **kw):
            if "独有说明5" in (caption or ""):
                return _parse_failed()
            return _good(caption)

        settings_sleep = lambda seconds: sleeps.append(seconds)        # noqa: E731
        from dataclasses import replace
        service = IngestService(
            replace(self.settings, description_retry_backoff_seconds=3.5),
            self.store, parser=None, index_service=IndexService(),
            chat_model_factory=lambda: self.base_model, progress_store=self.prog,
            retry_model_factory=self._factory, sleep=settings_sleep)
        import deephoto.pipeline.ingest as ingest_mod
        original = ingest_mod.describe_image
        ingest_mod.describe_image = fake
        try:
            service.ingest(self.doc_id)
        finally:
            ingest_mod.describe_image = original
        self.assertEqual(sleeps, [3.5])                                # 2 轮之间退避 1 次


@unittest.skipUnless(HAVE_DEPS, "需要 pillow")
class ParseDescriptionCodeTest(unittest.TestCase):
    """§4.5:各失败分支的机读 code;ok 返回结构不变(不含 reason/code/truncated)。"""

    def test_empty_output(self):
        r = parse_description_json("   ")
        self.assertEqual(r["diagnostic"]["code"], "empty")
        self.assertEqual(r["diagnostic"]["result"], "parse_failed")
        self.assertTrue(r["diagnostic"]["reason"])

    def test_no_json_object(self):
        r = parse_description_json("这是一段没有 JSON 的文字")
        self.assertEqual(r["diagnostic"]["code"], "no_json")

    def test_invalid_json(self):
        r = parse_description_json('{visible_summary: "缺引号"}')
        self.assertEqual(r["diagnostic"]["code"], "invalid_json")

    def test_field_type(self):
        r = parse_description_json('{"visible_summary": 123}')
        self.assertEqual(r["diagnostic"]["code"], "field_type")
        self.assertIn("visible_summary", r["diagnostic"]["reason"])    # 点名出错字段

    def test_ok_shape_unchanged(self):
        r = parse_description_json(
            '{"visible_summary": "a", "visible_labels": ["x"], "caption": "c",'
            ' "context_summary": "s", "uncertain_details": ["u"]}')
        self.assertEqual(r["diagnostic"], {"result": "ok"})            # 不多出机读键
        self.assertEqual(r["visible_summary"], "a")
        self.assertEqual(r["visible_labels"], ["x"])


if __name__ == "__main__":
    unittest.main()
