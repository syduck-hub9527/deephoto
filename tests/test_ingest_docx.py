"""docx 端到端入库:上传字段 -> 本地解析 -> 分块/图文关联 -> 索引 -> 检索(P2 §5 验收)。

不经网络、不经 MinerU;fixture 由 python-docx 程序生成。
"""

import io
import tempfile
import unittest
from pathlib import Path

import _bootstrap  # noqa: F401

try:
    import numpy  # noqa: F401
    from PIL import Image
    HAVE_DEPS = True
except ImportError:
    HAVE_DEPS = False

if HAVE_DEPS:
    from deephoto import repo
    from deephoto.agent.knowledge import KnowledgeService
    from deephoto.config import Settings
    from deephoto.db import connect, init_db
    from deephoto.indexing.service import IndexService
    from deephoto.pipeline.ingest import IngestService
    from deephoto.security import LOCAL_CTX
    from deephoto.storage import ObjectStore


def _png(color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buf, format="PNG")
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading("第1章 存储器", level=1)
    doc.add_paragraph("存储器容量持续提升,如图 1.1 所示,这份文档的正文写得足够长以越过回退阈值。")
    p = doc.add_paragraph()
    p.add_run().add_picture(io.BytesIO(_png()))
    doc.add_paragraph("图 1.1 容量与尺寸", style="Caption")
    doc.add_heading("1.1 工艺", level=2)
    t = doc.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text = "元素"
    t.rows[0].cells[1].text = "类型"
    t.rows[1].cells[0].text = "磷"
    t.rows[1].cells[1].text = "N 型"
    t.rows[2].cells[0].text = "硼"
    t.rows[2].cells[1].text = "P 型"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@unittest.skipUnless(HAVE_DEPS, "需要 numpy、pillow 与 python-docx")
class IngestDocxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        root = Path(self.tmp.name)
        self.settings = Settings(
            moonshot_api_key=None, moonshot_base_url="", chat_model="k3", chat_temperature=1.0,
            embedding_base_url=None, embedding_api_key=None, embedding_model=None,
            data_dir=root, max_upload_mb=100, ingestion_version="v2", mineru_api_key=None)
        init_db(self.settings.db_path)
        self.conn = connect(self.settings.db_path)
        self.store = ObjectStore(root / "objects")
        data = _docx_bytes()
        key, digest = self.store.put(data, "sources", format_by_key_mime(), ext="docx")
        self.doc_id = repo.insert_document(
            self.conn, tenant_id=LOCAL_CTX.tenant_id, owner_id=LOCAL_CTX.user_id,
            filename="教材.docx", source_object_key=key, sha256=digest, ingestion_version="v2",
            source_format="docx", locator_kind="section", parse_engine="local")
        self.index = IndexService()
        IngestService(self.settings, self.store, parser=None,
                      index_service=self.index).ingest(self.doc_id)
        self.knowledge = KnowledgeService(self.store, self.index)

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def test_ready_with_sections_and_table_text(self):
        doc = repo.get_document(self.conn, self.doc_id)
        self.assertEqual(doc["status"], "ready")
        chunks = repo.chunks_for_document(self.conn, self.doc_id)
        sections = {c["section"] for c in chunks}
        self.assertIn("第1章 存储器", sections)
        self.assertIn("第1章 存储器 > 1.1 工艺", sections)
        blob = "\n".join(c["text"] for c in chunks)
        self.assertIn("磷 | N 型", blob)          # 表格文本可检索

    def test_figure_persisted_with_caption(self):
        occs = repo.occurrences_for_document(self.conn, self.doc_id)
        self.assertEqual(len(occs), 1)
        self.assertEqual(occs[0]["figure_number"], "1.1")
        self.assertEqual(occs[0]["bbox_coord"], "none")
        asset = repo.get_asset(self.conn, occs[0]["image_asset_id"])
        self.assertEqual(asset["mime_type"], "image/png")
        self.assertTrue(self.store.get(asset["original_object_key"]).startswith(b"\x89PNG"))

    def test_search_gives_section_locator(self):
        result = self.knowledge.search(self.conn, LOCAL_CTX, "容量 尺寸 存储器")
        self.assertTrue(result["chunks"])
        for c in result["chunks"]:
            self.assertNotIn("page_start", c)
            self.assertIn("第1章 存储器", c["locator"])
            self.assertEqual(c["document"], "教材.docx")
        images = result["images"]
        self.assertTrue(any(i["figure_number"] == "1.1" for i in images))

    def test_dedup_engine_local_vs_mineru(self):
        digest = ObjectStore.sha256_of(_docx_bytes())
        hit = repo.find_ready_document_by_hash(
            self.conn, LOCAL_CTX.tenant_id, digest, "v2", "local")
        self.assertEqual(hit["id"], self.doc_id)
        self.assertIsNone(repo.find_ready_document_by_hash(
            self.conn, LOCAL_CTX.tenant_id, digest, "v2", "mineru"))


def format_by_key_mime() -> str:
    from deephoto.parsing.formats import format_by_key
    return format_by_key("docx").mime


if __name__ == "__main__":
    unittest.main()
