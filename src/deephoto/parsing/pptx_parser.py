"""PPT(.pptx)本地解析(python-pptx;DEV_multi_format_documents.md §3.4c)。

- 一张幻灯片 = 一个 ParsedPage(locator_kind='slide',页码即幻灯片序号,天然有位置);
- 标题占位符 -> 该页 section("幻灯片 N:标题");标题文字本身仍按阅读顺序落为段落
  (标题可检索,与 docx/md 的标题处理一致);
- 形状按"先上后下、同排先左后右"排阅读顺序(python-pptx 返回的是层叠顺序);
  同排判定按形状高度相对比较(纵向错位 ≤ 0.3 × 较矮形状高度),容忍手工拖放
  的微小错位;不用固定格子分桶(跨桶界仍交错、同桶内上下堆叠会按 left 排错);
- 文本框/占位符/自选形状的文字经 text_frame 提取;组合形状递归展开;
  页脚/页码/日期/页眉占位符跳过(版式套话,入索引会污染检索排序);
- 图片占位符(PlaceholderPicture,内置"图片+说明"版式)与普通图片同样取图;
  图片经 a:blip 的 r:embed 直取图片 part 字节(含 content_type,不经格式识别);
  EMF/WMF/SVG 跳过并告警(与 docx 同约定),其余格式由入库处统一转 PNG;
- 图注:阅读顺序上图片之后第一条 match_caption 命中、位置在图片下方、水平重叠
  且距图片底边不超过 1 英寸的未配对文本(跳过普通段落与同排的图;遇下一行的
  图/表格/位置不符的图注即停);配对后不再单独出文本段落;
  位置信息缺失时退化为纯邻接判断;
- 表格按行展平(" | " 连接),复用 content_list._table_paragraphs(续段重复表头);
  被合并的格子(is_spanned)跳过,不重复输出;
- 讲者备注:该页最后一个段落,前缀"备注:";
- 连接线/直线是纯装饰,静默跳过;图表/SmartArt/OLE 等确有内容但读不到的形状
  记告警并带幻灯片页号;整篇无任何文字与图时由 validate_parsed 报可读错误。
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass

from ..config import Settings
from .base import (
    KIND_EMBEDDED_BITMAP,
    ParsedCaption,
    ParsedDocument,
    ParsedFigure,
    ParsedPage,
    ParsedParagraph,
)
from .captions import match_caption
from .content_list import PAGE_H, PAGE_W, _table_paragraphs
from .docx_parser import _NON_RASTER_MIMES
from .formats import check_zip_safety
from .registry import SourceFile

logger = logging.getLogger(__name__)


@dataclass
class _Item:
    """一页幻灯片上按阅读顺序展平的内容块。"""
    kind: str                      # text | image | table
    shape: object                  # 来源形状(图注位置判断用)
    text: str = ""                 # text/table: 文本内容
    image_bytes: bytes | None = None


_ROW_HEIGHT_RATIO = 0.3      # 同排判定:纵向错位不超过较矮形状高度的 0.3 倍(估值)
_ROW_FALLBACK_EMU = 274320   # 高度缺失时的固定同排容差(0.3 英寸)
# 图注顶边距图片底边的最大距离(1 英寸):图注应紧邻图片下方;无上限会把远处
# 以"图 N"开头的正文误配成图注(错误关联的置信度却是 1.0)
_CAPTION_MAX_GAP_EMU = 914400


def _reading_order(shapes) -> list:
    """阅读顺序:先上后下、同排先左后右。

    按 top 排序后贪心分行:与当前行的锚形状(行内最上方)比较,纵向错位
    ≤ 0.3 × min(两者高度) 归入同一排,否则另起一排;排内按 left。
    不用固定格子分桶:错位恰跨格子边界仍会交错;同格内上下堆叠的小文本框
    会被按 left 把下方框排到上方框前面。位置缺失按 0;高度缺失退回固定容差。
    """
    def top_of(shape) -> int:
        return shape.top if shape.top is not None else 0

    def left_of(shape) -> int:
        return shape.left if shape.left is not None else 0

    rows: list[tuple[object, list]] = []     # (锚形状, 行成员)
    for shape in sorted(shapes, key=lambda sh: (top_of(sh), left_of(sh))):
        if rows:
            anchor, members = rows[-1]
            if anchor.height is not None and shape.height is not None:
                limit = _ROW_HEIGHT_RATIO * min(anchor.height, shape.height)
            else:
                limit = _ROW_FALLBACK_EMU
            if top_of(shape) - top_of(anchor) <= limit:
                members.append(shape)
                continue
        rows.append((shape, [shape]))
    out: list = []
    for _, members in rows:
        out.extend(sorted(members, key=left_of))
    return out


def _is_chrome_placeholder(shape) -> bool:
    """页脚/页码/日期/页眉占位符:版式套话不是正文,跳过——否则每页多出几条无意义
    文本("1"、"2026-09-29"、公司页脚)进索引,页脚文字还会在几乎每个块里命中,
    污染 BM25 排序。标题/正文/图片占位符不在此列。"""
    if not getattr(shape, "is_placeholder", False):
        return False
    from pptx.enum.shapes import PP_PLACEHOLDER
    return shape.placeholder_format.type in {
        PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER,
        PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER,
    }


def _caption_position_ok(img, cap) -> bool:
    """图注位置约束:文本框在图片下方(顶边不低于图片顶边)、水平区间有重叠、
    且顶边距图片底边不超过 1 英寸(图注应紧邻图片下方,无距离上限会把远处
    以"图 N"开头的正文误配为图注)。任一位置信息缺失时放行(退化为纯邻接)。"""
    values = (img.top, img.left, img.width, cap.top, cap.left, cap.width)
    if any(v is None for v in values):
        return True
    if cap.top < img.top:
        return False
    if not (cap.left < img.left + img.width and img.left < cap.left + cap.width):
        return False
    bottom = img.top + (img.height if img.height is not None else 0)
    return cap.top - bottom <= _CAPTION_MAX_GAP_EMU


def _table_text(table) -> str:
    """表格按行展平为文本行;被合并的格子(is_spanned,独立 tc 但无文字)跳过。"""
    lines: list[str] = []
    for row in table.rows:
        texts = [cell.text.strip() for cell in row.cells if not cell.is_spanned]
        row_text = " | ".join(t for t in texts if t)
        if row_text.strip(" |"):
            lines.append(row_text)
    return "\n".join(lines)


class PptxParser:
    """pptx -> ParsedDocument(locator_kind='slide',一张幻灯片一页)。纯本地,无网络。"""

    engine = "local"

    def __init__(self, settings: Settings):
        self._settings = settings

    def parse(self, src: SourceFile, observer):
        from pptx import Presentation

        check_zip_safety(src.data, self._settings.max_zip_uncompressed_mb)
        prs = Presentation(io.BytesIO(src.data))

        pages: list[ParsedPage] = []
        counters = {"para": 0, "cap": 0, "fig": 0}
        placeholder = (0.0, 0.0, PAGE_W, PAGE_H)

        def add_paragraph(page: ParsedPage, section: str, text: str) -> None:
            text = text.strip()
            if not text:
                return
            counters["para"] += 1
            page.paragraphs.append(ParsedParagraph(
                id=f"p{page.page_number}_c{counters['para']}", text=text,
                page_number=page.page_number, bbox=placeholder, section=section,
            ))

        for number, slide in enumerate(prs.slides, 1):
            page = ParsedPage(page_number=number, width=PAGE_W, height=PAGE_H, is_scanned=False)
            title_shape = slide.shapes.title
            title = (title_shape.text or "").strip() if title_shape is not None else ""
            section = f"幻灯片 {number}:{title}" if title else f"幻灯片 {number}"

            items = self._walk_shapes(slide.shapes, observer, number)
            consumed: set[int] = set()      # 已配对为图注的文本 item 下标
            for i, item in enumerate(items):
                if i in consumed:
                    continue
                if item.kind == "text":
                    add_paragraph(page, section, item.text)
                elif item.kind == "table":
                    for piece in _table_paragraphs("", item.text, ""):
                        add_paragraph(page, section, piece)
                else:                        # image
                    caption: str | None = None
                    j = _find_caption(items, i, consumed, item.shape)
                    if j is not None:
                        caption = items[j].text
                        consumed.add(j)
                    caption_id = None
                    if caption:
                        counters["cap"] += 1
                        caption_id = f"p{number}_cap{counters['cap']}"
                        num, caption_text = match_caption(caption)
                        page.captions.append(ParsedCaption(
                            id=caption_id, text=caption_text, figure_number=num,
                            page_number=number, bbox=placeholder,
                        ))
                    # 图注文本必须落进段落(图文关联按"图注文本出现在块内"判定 caption_of)
                    add_paragraph(page, section, caption or "")
                    counters["fig"] += 1
                    page.figures.append(ParsedFigure(
                        id=f"p{number}_fig{counters['fig']}", page_number=number,
                        bbox=placeholder, kind=KIND_EMBEDDED_BITMAP,
                        image_bytes=item.image_bytes, caption_id=caption_id,
                    ))

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    add_paragraph(page, section, f"备注:{notes}")
            pages.append(page)

        return ParsedDocument(page_count=len(pages), pages=pages, locator_kind="slide")

    def _walk_shapes(self, shapes, observer, page_number: int) -> list[_Item]:
        """形状按阅读顺序展平:文本、图片、表格;组合形状递归(组内同样按阅读顺序)。"""
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        out: list[_Item] = []
        for shape in _reading_order(shapes):
            if _is_chrome_placeholder(shape):
                continue
            try:
                shape_type = shape.shape_type
            except NotImplementedError:
                shape_type = None           # 无预设几何等未识别形状:按文本/表格兜底
            if shape_type == MSO_SHAPE_TYPE.GROUP:
                out.extend(self._walk_shapes(shape.shapes, observer, page_number))
            # 图片占位符(PlaceholderPicture)shape_type 是 PLACEHOLDER 而非 PICTURE,
            # 但同样有 _pic;只认 PICTURE 会静默丢图(内置"图片+说明"版式很常见)
            elif shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "_pic"):
                blob = self._image_bytes(shape, observer, page_number)
                if blob:
                    out.append(_Item("image", shape, image_bytes=blob))
            elif getattr(shape, "has_table", False) and shape.has_table:
                body = _table_text(shape.table)
                if body:
                    out.append(_Item("table", shape, text=body))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    out.append(_Item("text", shape, text=text))
                # 空文本的形状(装饰用矩形等)极常见,不告警
            elif shape_type == MSO_SHAPE_TYPE.LINE:
                continue                    # 连接线/直线是纯装饰:没有可入库内容,不告警
            else:
                # 图表/SmartArt/OLE 等确有内容但读不到的形状:记告警(带页号),不再静默
                label = str(shape_type) if shape_type is not None else "未识别"
                observer.warn(f"幻灯片 {page_number}:跳过暂不支持的形状({label}),其内容未入库")
        return out

    def _image_bytes(self, shape, observer, page_number: int) -> bytes | None:
        """图片字节:经 a:blip 的 r:embed 直取图片 part(与 docx 同法,不经 python-pptx 的
        格式识别,未知格式不会炸);链接图(r:link,无嵌入字节)跳过;
        EMF/WMF/SVG 跳过并告警(Pillow 支持依平台而异,与 docx 同约定)。"""
        from pptx.oxml.ns import qn

        blip = shape._pic.find(".//" + qn("a:blip"))
        rid = blip.get(qn("r:embed")) if blip is not None else None
        if not rid:
            return None
        part = shape.part.related_part(rid)
        content_type = (getattr(part, "content_type", "") or "").lower()
        if content_type in _NON_RASTER_MIMES:
            observer.warn(f"幻灯片 {page_number}:跳过暂不支持的图片格式({content_type}),其余内容不受影响")
            return None
        return part.blob


def _find_caption(items: list[_Item], after: int, consumed: set[int], img) -> int | None:
    """图片的图注:阅读顺序上其后第一条 match_caption 命中、位于图片下方且水平重叠
    的未配对文本。扫描规则:
    - 跳过普通文本段落("图片+说明"版式的标题就位于图片与图注之间);
    - 跳过同排/上方的图片(并排双图各自配对,互不抢注);
    - 遇到明确位于图片下方的下一张图、或表格:停止(不抢下一行内容的图注);
    - 遇到图注样式但位置不符的文本:停止(那是别人的图注,不越过去找更远的)。
    """
    bottom = None
    if img.top is not None and img.height is not None:
        bottom = img.top + img.height
    for j in range(after + 1, len(items)):
        if j in consumed:
            continue
        item = items[j]
        if item.kind == "image":
            top = item.shape.top
            if bottom is not None and top is not None and top >= bottom:
                return None                    # 下一行的图:本图的图注不会再出现
            continue
        if item.kind == "table":
            return None
        if match_caption(item.text)[0] is None:
            continue                           # 普通段落:跳过,继续找
        if _caption_position_ok(img, item.shape):
            return j
        return None                            # 图注样式但位置不符:是别人的图注
    return None
